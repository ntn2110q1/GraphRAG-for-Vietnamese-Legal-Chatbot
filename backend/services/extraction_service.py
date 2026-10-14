import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument

from backend.core.config import settings


ARTICLE_HEADING = re.compile(r"(?im)^\s*Điều\s+(\d+[a-zA-Z]?)\s*\.\s*(.*)$")


@dataclass
class TextChunk:
    index: int
    content: str
    article_number: str | None
    location_label: str


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text or "")
    text = text.replace("\x00", " ").replace("\ufeff", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]*\n(?:[ \t]*\n)+", "\n\n", text)
    return text.strip()


def extract_file_text(path: Path) -> str:
    extension = path.suffix.casefold()
    if extension == ".txt":
        raw = path.read_text(encoding="utf-8-sig", errors="replace")
    elif extension == ".docx":
        doc = DocxDocument(path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        raw = "\n".join(paragraphs)
    elif extension == ".pdf":
        with pymupdf.open(path) as pdf:
            raw = "\n\n".join(page.get_text("text") for page in pdf)
        if len(raw.strip()) < 80:
            raise ValueError("PDF has no usable text layer. Configure OCR for scanned PDFs and retry.")
    else:
        raise ValueError("Supported file types are PDF, DOCX, and TXT.")
    cleaned = normalize_text(raw)
    if len(cleaned) < 20:
        raise ValueError("The extracted document is empty or too short to index.")
    return cleaned


def _split_long_section(section: str, max_chars: int, overlap: int) -> list[str]:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", section) if p.strip()]
    parts: list[str] = []
    current = ""
    for paragraph in paragraphs:
        while len(paragraph) > max_chars:
            if current:
                parts.append(current)
                current = ""
            parts.append(paragraph[:max_chars].strip())
            paragraph = paragraph[max(1, max_chars - overlap):].strip()
        if not current:
            current = paragraph
        elif len(current) + len(paragraph) + 2 <= max_chars:
            current += "\n\n" + paragraph
        else:
            parts.append(current)
            tail = current[-overlap:] if overlap else ""
            current = (tail + "\n\n" + paragraph).strip() if tail else paragraph
    if current:
        parts.append(current)
    return parts


def make_chunks(text: str) -> list[TextChunk]:
    max_chars = max(300, settings.legal_chunk_size)
    overlap = max(0, min(settings.legal_chunk_overlap, max_chars // 3))
    headings = list(ARTICLE_HEADING.finditer(text))
    sections: list[tuple[str | None, str, str]] = []
    if headings:
        if headings[0].start() > 0 and text[:headings[0].start()].strip():
            sections.append((None, "Phần mở đầu", text[:headings[0].start()].strip()))
        for idx, match in enumerate(headings):
            stop = headings[idx + 1].start() if idx + 1 < len(headings) else len(text)
            number = match.group(1)
            heading_title = match.group(2).strip()
            label = f"Điều {number}" + (f". {heading_title}" if heading_title else "")
            sections.append((number, label, text[match.start():stop].strip()))
    else:
        sections.append((None, "Toàn văn", text))

    chunks: list[TextChunk] = []
    for article_number, label, section in sections:
        for part in _split_long_section(section, max_chars, overlap):
            if len(part.strip()) >= 20:
                chunks.append(TextChunk(len(chunks), part, article_number, label))
    return chunks
